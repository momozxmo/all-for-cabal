"""Direct Aztek Item ID imports must never lose a bundle or touch Aztek."""
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook


HEADERS = ['Bundle Name', 'Item ID', 'Qty', 'Tier', 'Item Name']


def workbook_bytes(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows, hidden in sheets:
        sheet = book.create_sheet(name)
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        if hidden:
            sheet.sheet_state = 'hidden'
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def upload(client, sheets):
    return client.post('/api/bundles/import', files={
        'file': ('bundles.xlsx', workbook_bytes(sheets),
                 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})


def test_groups_in_source_order_and_keeps_same_id_in_different_bundles(client):
    response = upload(client, [('ชุด ก', [
        ['Pack A', '200479', 3, 'Epic', 'Potion'],
        ['Pack B', '200479', 5, None, None],
        ['Pack A', 200480, None, None, 'Orb'],
    ], False), ('ชุด ข', [['Pack A', '33', 2, 'Rare']], False),
        ('old', [['Wrong', '99', 1]], True)])
    assert response.status_code == 200
    data = response.json()
    assert [s['name'] for s in data['sheets']] == ['ชุด ก', 'ชุด ข']
    bundles = data['sheets'][0]['bundles']
    assert [b['name'] for b in bundles] == ['Pack A', 'Pack B']
    assert [(i['id'], i['qty'], i['tier']) for i in bundles[0]['items']] == [
        ('200479', '3', 'Epic'), ('200480', '1', 'Common')]
    assert bundles[1]['items'][0]['id'] == '200479'
    assert bundles[1]['items'][0]['qty'] == '5'
    assert bundles[0]['items'][0]['name'] == 'Potion'
    assert data['sheets'][1]['bundles'][0]['items'][0]['id'] == '33'


@pytest.mark.parametrize('row,field', [
    (['Pack', '', 1], 'Item ID'), (['Pack', 1.5, 1], 'Item ID'),
    (['Pack', True, 1], 'Item ID'), (['Pack', '=1+1', 1], 'สูตร'),
    (['Pack', '9999999999999', 1], 'Item ID'),
    (['Pack', 1, 0], 'Qty'), (['Pack', 1, -1], 'Qty'),
    (['Pack', 1, 1, 'wrong'], 'Tier'), (['', 1, 1], 'Bundle Name'),
])
def test_invalid_row_blocks_whole_sheet_with_excel_row_number(client, row, field):
    response = upload(client, [('Data', [['Valid', '11', 1], row], False)])
    assert response.status_code == 200
    sheet = response.json()['sheets'][0]
    assert sheet['bundles'] == []
    assert sheet['errors'][0]['row'] == 3
    assert field in sheet['errors'][0]['message']


def test_duplicate_within_bundle_is_not_silently_dropped(client):
    response = upload(client, [('Data', [['Pack', 11, 1], ['Pack', 11, 3]], False)])
    sheet = response.json()['sheets'][0]
    assert sheet['bundles'] == []
    assert sheet['errors'][0]['row'] == 3
    assert 'ซ้ำ' in sheet['errors'][0]['message']


def test_import_requires_authentication(anonymous_client):
    assert upload(anonymous_client, [('Data', [], False)]).status_code == 401


def test_bad_workbook_returns_readable_error(client):
    response = client.post('/api/bundles/import', files={'file': ('bad.xlsx', b'bad')})
    assert response.status_code == 400
    assert 'Excel' in response.json()['detail']


def test_download_template_is_an_attachment_without_real_sample_ids(client):
    response = client.get('/api/bundles/template')
    assert response.status_code == 200
    assert 'attachment' in response.headers['content-disposition']
    book = load_workbook(BytesIO(response.content))
    sheet = book.active
    assert sheet['A2'].value == '[Item ID 1]'
    book.close()
    result = client.post('/api/bundles/import', files={'file': ('blank.xlsx', response.content)})
    assert result.status_code == 200
    assert result.json()['sheets'][0]['bundles'] == []


def test_bundle_upload_uses_workbook_size_limit():
    from web.request_limits import request_limit, WORKBOOK_BODY_MAX
    assert request_limit('POST', '/api/bundles/import') == WORKBOOK_BODY_MAX


def test_bad_header_does_not_hide_other_visible_sheets(client):
    data = workbook_bytes([('Bad', [], False), ('Good', [['Pack', 123, 1]], False)])
    book = load_workbook(BytesIO(data))
    book['Bad']['A1'] = 'Wrong header'
    out = BytesIO()
    book.save(out)
    book.close()
    response = client.post('/api/bundles/import', files={'file': ('test.xlsx', out.getvalue())})
    sheets = response.json()['sheets']
    assert [s['name'] for s in sheets] == ['Bad', 'Good']
    assert sheets[0]['errors'][0]['row'] == 1
    assert sheets[1]['bundles'][0]['items'][0]['id'] == '123'


def test_excess_items_are_not_silently_truncated(client):
    response = upload(client, [('Data', [['Pack', i, 1] for i in range(1, 202)], False)])
    sheet = response.json()['sheets'][0]
    assert sheet['bundles'] == []
    assert sheet['errors'][0]['row'] == 202


def test_preview_script_is_served_as_javascript(client):
    response = client.get('/static/bundle_import.js')
    assert response.status_code == 200
    assert 'javascript' in response.headers['content-type']


def test_input_workbook_is_not_modified(tmp_path):
    from web.bundle_import import read_bundle_template
    raw = workbook_bytes([('Data', [['Pack', 123, 1]], False)])
    source = tmp_path / 'test.xlsx'
    source.write_bytes(raw)
    read_bundle_template(source)
    assert source.read_bytes() == raw


def test_stale_sheet_dimensions_do_not_drop_rows(client):
    import re
    from zipfile import ZipFile
    raw = workbook_bytes([('Data', [['A', 11, 1], ['B', 22, 1]], False)])
    out = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(out, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(b'<dimension ref="[^"]+"', b'<dimension ref="A1:E2"', data)
            target.writestr(entry, data)
    response = client.post('/api/bundles/import', files={'file': ('test.xlsx', out.getvalue())})
    assert [b['name'] for b in response.json()['sheets'][0]['bundles']] == ['A', 'B']
