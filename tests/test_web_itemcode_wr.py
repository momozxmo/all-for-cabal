# -*- coding: utf-8 -*-
"""Observable Mastercode WR import behavior at the Item Code API seam."""
import io

import openpyxl


MTH = 'CabalM TH'
XLSX = ('application/vnd.openxmlformats-officedocument.'
        'spreadsheetml.sheet')


def _add_wr_block(sheet, *, row, column=2, header, expires,
                  game='CABAL M TH', mastercode='CODE001',
                  bundle_id=224728, usage_limit=200, fill=None):
    header_cell = sheet.cell(row=row, column=column, value=header)
    if fill:
        header_cell.fill = openpyxl.styles.PatternFill(
            'solid', fgColor=fill)
    sheet.cell(row=row + 1, column=column, value='CODE EXPIRE DATE')
    sheet.cell(row=row + 1, column=column + 2, value=expires)
    sheet.cell(row=row + 3, column=column, value='Game')
    sheet.cell(row=row + 3, column=column + 2, value=game)
    sheet.cell(row=row + 3, column=column + 6,
               value='Limit การใช้งาน')
    sheet.cell(row=row + 3, column=column + 10, value=usage_limit)
    sheet.cell(row=row + 4, column=column, value='Mastercode')
    sheet.merge_cells(
        start_row=row + 4, start_column=column,
        end_row=row + 4, end_column=column + 1)
    sheet.cell(row=row + 4, column=column + 2, value=mastercode)
    sheet.merge_cells(
        start_row=row + 4, start_column=column + 2,
        end_row=row + 4, end_column=column + 4)
    sheet.cell(row=row + 4, column=column + 5, value=bundle_id)


def _save_book(book):
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _wr_workbook_bytes(*, sheet_name='12.07 M'):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = sheet_name
    blocks = [
        (1, 'Master Code 12/07/26 - Code 1 (เจ็นเองได้เลย)',
         'ALPHA001', 224728, 200),
        (20, 'Master Code 12/07/26 - Code 2 (หมายเหตุ)',
         'BETA002', 224729, 200),
        (38, 'Master Code 12/07/26 - Code 3',
         'GAMMA003', 224730, 400),
    ]
    for row, header, mastercode, bundle_id, limit in blocks:
        _add_wr_block(
            sheet, row=row, header=header,
            expires='12/07/26, เวลา 23.59 น.', mastercode=mastercode,
            bundle_id=bundle_id, usage_limit=limit)
    return _save_book(book)


def _upload_wr(client, payload, *, game=MTH):
    response = client.post(
        '/api/itemcodes/import',
        data={'game': game, 'import_mode': 'mastercode_wr'},
        files={'file': ('wr.xlsx', payload, XLSX)},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _apply_wr(client, pending_id, selected_sheets, *, game=MTH):
    response = client.post('/api/itemcodes/import/apply', json={
        'pending_id': pending_id,
        'selected_sheets': selected_sheets,
        'game': game,
    })
    assert response.status_code == 200, response.text
    return response.json()['preview_rows']


def test_mastercode_wr_happy_path_upload_select_and_preview(client):
    uploaded = client.post(
        '/api/itemcodes/import',
        data={'game': MTH, 'import_mode': 'mastercode_wr'},
        files={'file': ('wr.xlsx', _wr_workbook_bytes(), XLSX)},
    )

    assert uploaded.status_code == 200, uploaded.text
    started = uploaded.json()
    assert started['import_mode'] == 'mastercode_wr'
    assert started['pending_id']
    assert started['sheets'] == [{
        'name': '12.07 M',
        'display_name': '12.07 M',
        'count': 3,
    }]
    assert 'preview_rows' not in started

    applied = client.post('/api/itemcodes/import/apply', json={
        'pending_id': started['pending_id'],
        'selected_sheets': ['12.07 M'],
        'game': MTH,
    })

    assert applied.status_code == 200, applied.text
    body = applied.json()
    rows = body['preview_rows']
    assert [row['block'] for row in rows] == ['Code 1', 'Code 2', 'Code 3']
    assert [row['name'] for row in rows] == [
        '1. Master Code 12/07/26 - Code 1 - ALPHA001',
        '2. Master Code 12/07/26 - Code 2 - BETA002',
        '3. Master Code 12/07/26 - Code 3 - GAMMA003',
    ]
    assert rows[0]['sheet'] == '12.07 M'
    assert rows[0]['game'] == MTH
    assert rows[0]['start_time'] == '2026-07-12T00:00:00'
    assert rows[0]['end_time'] == '2026-07-12T23:59:59'
    assert rows[0]['mastercode'] == 'ALPHA001'
    assert rows[0]['bundle_id'] == '224728'
    assert rows[0]['usage_limit'] == '200'
    assert rows[0]['ready'] is True
    assert rows[0]['selectable'] is True
    assert rows[0]['selected'] is True

    draft = rows[0]['draft']
    assert draft['name_th'] == rows[0]['name']
    assert draft['name_en'] == rows[0]['name']
    assert draft['slug'] == (
        '1-master-code-120726-code-1-alpha001-mth')
    assert draft['uses_per_user'] == '1'
    assert draft['limited'] is True
    assert draft['quantity'] == draft['remaining'] == '200'
    assert draft['start_time'] == rows[0]['start_time']
    assert draft['end_time'] == rows[0]['end_time']
    assert draft['rewards'] == [{
        'name_th': rows[0]['name'],
        'name_en': rows[0]['name'],
        'uses_per_user': '1',
        'limited': True,
        'quantity': '200',
        'remaining': '200',
        'code_type': '1',
        'code_list': 'ALPHA001',
        'prefix': '',
        'num_codes': '',
        'bundle_id': '224728',
    }]


def test_mastercode_wr_orders_horizontal_blocks_and_keeps_daily_sequence(
        client):
    book = openpyxl.Workbook()
    first = book.active
    first.title = ' Leading Daily'
    _add_wr_block(
        first, row=1, column=2,
        header='Master Code 26/08/26 - Code 1 (แดง)',
        expires='27/08/26, เวลา 23.59 น.', mastercode='FIRST',
        fill='FF0000')
    _add_wr_block(
        first, row=1, column=14,
        header='Master Code 27/08/26 - Code 2 (ส้ม)',
        expires='27/08/26, เวลา 23.59 น.', mastercode='SECOND',
        bundle_id=224729, fill='F4B183')
    _add_wr_block(
        first, row=20, column=2,
        header='Master Code 27/08/26 - Code 1 (เหลือง)',
        expires='27/08/26, เวลา 23.59 น.', game='CABAL PC TH',
        mastercode='PCFIRST', bundle_id=224730, fill='FFFF00')
    _add_wr_block(
        first, row=20, column=14,
        header='Master Code 31/08/26 - Code 9',
        expires='31/02/26, เวลา 23.59 น.', game='CABAL UNKNOWN',
        mastercode='UNKNOWN', bundle_id=224731)
    first.cell(
        row=50, column=2,
        value='Master Code 01/09/26 - Code 99 (เป็นเพียงข้อความอ้างอิง)')
    second = book.create_sheet('Second Daily')
    _add_wr_block(
        second, row=1,
        header='Master Code 27/08/26 - Code 1 (อีกชีท)',
        expires='27/08/26, เวลา 23.59 น.', mastercode='THIRD',
        bundle_id=224732)
    payload = _save_book(book)

    selected_only = _upload_wr(client, payload)
    assert [sheet['name'] for sheet in selected_only['sheets']] == [
        ' Leading Daily', 'Second Daily']
    only_second = _apply_wr(
        client, selected_only['pending_id'], ['Second Daily'])
    assert only_second[0]['daily_sequence'] == 3
    assert only_second[0]['name'].startswith(
        '3. Master Code 27/08/26 - Code 1 - THIRD')
    assert any('Same-Day Game Overlap' in warning
               for warning in only_second[0]['warnings'])

    all_sheets = _upload_wr(client, payload)
    rows = _apply_wr(
        client, all_sheets['pending_id'],
        [' Leading Daily', 'Second Daily'])
    assert [(row['sheet'], row['row'], row['column']) for row in rows] == [
        (' Leading Daily', 1, 2),
        (' Leading Daily', 1, 14),
        (' Leading Daily', 20, 2),
        (' Leading Daily', 20, 14),
        (' Leading Daily', 50, 2),
        ('Second Daily', 1, 2),
    ]
    assert [rows[index]['daily_sequence'] for index in (0, 1, 5)] == [1, 2, 3]
    assert rows[0]['date'] == '2026-08-27'
    assert rows[0]['base_header'] == 'Master Code 26/08/26 - Code 1'
    assert rows[2]['game'] == 'CabalPC TH'
    assert rows[2]['daily_sequence'] == 1
    assert rows[2]['selectable'] is False
    assert 'CabalPC TH' in rows[2]['disabled_reason']
    assert rows[3]['date'] == ''
    assert rows[3]['game'] == ''
    assert rows[3]['ready'] is False
    assert rows[3]['selectable'] is False
    assert any('CODE EXPIRE DATE' in issue for issue in rows[3]['issues'])
    assert any('ไม่รู้จัก Game' in issue for issue in rows[3]['issues'])
    assert rows[4]['block'] == 'Code 99'
    assert rows[4]['ready'] is False
    assert rows[4]['row'] == 50
    assert rows[4]['issues']


def test_mastercode_wr_marks_incomplete_rows_but_keeps_them_as_drafts(client):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Incomplete'
    _add_wr_block(
        sheet, row=1, header='Master Code 28/08/26 - Code 1',
        expires='28/08/26, เวลา 23.59 น.', mastercode=None,
        bundle_id=224728.0, usage_limit=200)
    _add_wr_block(
        sheet, row=20, header='Master Code 28/08/26 - Code 2',
        expires='28/08/26, เวลา 23.59 น.', mastercode='ANY-LENGTH-CODE',
        bundle_id=None, usage_limit=0)
    _add_wr_block(
        sheet, row=38, header='Master Code 28/08/26 - Code 3',
        expires='28/08/26, เวลา 23.59 น.', mastercode='ANY-LENGTH-CODE',
        bundle_id=224728, usage_limit=400)

    started = _upload_wr(client, _save_book(book))
    rows = _apply_wr(client, started['pending_id'], ['Incomplete'])

    assert rows[0]['mastercode'] == ''
    assert rows[0]['bundle_id'] == '224728'
    assert '[ยังไม่มี Mastercode]' in rows[0]['name']
    assert rows[0]['ready'] is False
    assert rows[0]['selectable'] is True
    assert rows[0]['selected'] is False
    assert rows[0]['draft']['rewards'][0]['code_list'] == ''
    assert rows[1]['mastercode'] == 'ANY-LENGTH-CODE'
    assert rows[1]['bundle_id'] == ''
    assert rows[1]['usage_limit'] == ''
    assert rows[1]['selected'] is False
    assert rows[2]['ready'] is True
    assert rows[2]['usage_limit'] == '400'
    assert rows[2]['draft']['quantity'] == '400'
    assert rows[2]['draft']['remaining'] == '400'


def test_mastercode_wr_pending_import_is_owner_scoped(
        client, client_for, other_member):
    started = _upload_wr(client, _wr_workbook_bytes())
    other_client = client_for(other_member)

    denied = other_client.post('/api/itemcodes/import/apply', json={
        'pending_id': started['pending_id'],
        'selected_sheets': ['12.07 M'],
        'game': MTH,
    })

    assert denied.status_code == 404
    assert _apply_wr(client, started['pending_id'], ['12.07 M'])


def test_mastercode_wr_exact_sheet_rejection_does_not_consume_pending(client):
    started = _upload_wr(client, _wr_workbook_bytes())

    rejected = client.post('/api/itemcodes/import/apply', json={
        'pending_id': started['pending_id'],
        'selected_sheets': ['12.07 M '],
        'game': MTH,
    })

    assert rejected.status_code == 400
    assert 'ไม่พบ sheet' in rejected.json()['detail']
    assert _apply_wr(client, started['pending_id'], ['12.07 M'])


def test_mastercode_wr_source_identity_is_stable_across_same_file_reimport(
        client):
    payload = _wr_workbook_bytes()
    first = _upload_wr(client, payload)
    first_source = _apply_wr(
        client, first['pending_id'], ['12.07 M'])[0]['source_id']
    second = _upload_wr(client, payload)
    second_source = _apply_wr(
        client, second['pending_id'], ['12.07 M'])[0]['source_id']

    assert first_source == second_source


def test_mastercode_wr_drafts_use_existing_run_validation_before_aztek(client):
    valid_started = _upload_wr(client, _wr_workbook_bytes())
    valid = _apply_wr(
        client, valid_started['pending_id'], ['12.07 M'])[0]['draft']

    valid_response = client.post('/api/itemcodes/run', json={
        'game': MTH, 'itemcodes': [valid], 'do_save': False,
    })

    # Reaching the existing Aztek-session gate proves the WR draft passed the
    # same public run contract; the test intentionally does not connect Aztek.
    assert valid_response.status_code == 409
    assert 'Aztek' in valid_response.json()['detail']

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Missing code'
    _add_wr_block(
        sheet, row=1, header='Master Code 28/08/26 - Code 1',
        expires='28/08/26, เวลา 23.59 น.', mastercode=None,
        bundle_id=224728, usage_limit=200)
    incomplete_started = _upload_wr(client, _save_book(book))
    incomplete = _apply_wr(
        client, incomplete_started['pending_id'], ['Missing code'])[0]['draft']

    incomplete_response = client.post('/api/itemcodes/run', json={
        'game': MTH, 'itemcodes': [incomplete], 'do_save': False,
    })

    assert incomplete_response.status_code == 400
    assert 'Code' in incomplete_response.json()['detail']


def test_mastercode_wr_requires_usage_limit_at_itemcode_and_reward(client):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Missing limit'
    _add_wr_block(
        sheet, row=1, header='Master Code 28/08/26 - Code 1',
        expires='28/08/26, เวลา 23.59 น.', mastercode='LIMIT001',
        bundle_id=224728, usage_limit=None)
    started = _upload_wr(client, _save_book(book))
    draft = _apply_wr(
        client, started['pending_id'], ['Missing limit'])[0]['draft']

    itemcode_response = client.post('/api/itemcodes/run', json={
        'game': MTH, 'itemcodes': [draft], 'do_save': False,
    })
    assert itemcode_response.status_code == 400
    assert 'Usage Limit' in itemcode_response.json()['detail']

    draft['limited'] = True
    draft['quantity'] = '200'
    draft['remaining'] = '200'
    reward_response = client.post('/api/itemcodes/run', json={
        'game': MTH, 'itemcodes': [draft], 'do_save': False,
    })
    assert reward_response.status_code == 400
    assert 'Usage Limit' in reward_response.json()['detail']
