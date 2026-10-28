"""Whole-feature review: exercise the real import/storage/handoff consumers."""
from copy import deepcopy
from io import BytesIO

from openpyxl import Workbook
from playwright.sync_api import expect
import pytest

from web.models import Job, WorkspaceRecord
from web.workspaces import WorkspaceRepository
from test_new_item_import_api import one_draft, preview, apply_preview
from test_new_item_reference_api import confirm
from test_new_item_create_flow import GAME, boundary, prepare, run
from test_new_item_reader import ORIGIN, saved_html
from test_new_item_recovery_api import reopen
from test_new_item_ui import operator_page
from test_web_validation import _connect_aztek


def pass_drafts(client, test_database, monkeypatch, drafts):
    """Seed prior saved identities; real public readback must pass before handoff."""
    _connect_aztek(client)
    selected, occurrences = [], []
    for index, draft in enumerate(drafts, 91):
        edited = client.patch('/api/new-items/drafts/' + draft['draft_id'], json={
            'revision': draft['revision'], 'values': {},
            'confirmed_fields': ['item_type', 'tradeable']}).json()
        frozen = confirm(client, edited, game=GAME)
        with test_database.session() as db:
            job = db.get(Job, draft['draft_id'])
            job.status = 'saved'
            job.result = {'creation': {'state': 'saved', 'item_id': str(index), 'evidence': {
                'detail_url': f'{ORIGIN}/combo/cabalm/shop/items/{index}'}}}
        html = saved_html(GAME).replace('Synthetic name', frozen['submitted_values']['name']).replace(
            'Synthetic description', frozen['submitted_values']['description'] or '').replace(
            '<span>91</span>', f'<span>{index}</span>')
        observed = boundary(monkeypatch, test_database, draft['draft_id'], saved=html, saved_id=index)
        checked = client.post('/api/new-items/' + draft['draft_id'] + '/recheck', json={})
        assert checked.status_code == 200 and checked.json()['recheck']['outcome'] == 'passed', checked.text
        assert not observed['posts'] and not observed['unexpected']
        selected.append(draft['draft_id'])
        occurrences.extend(o['occurrence_id'] for o in frozen['occurrences'])
    return {'game': GAME, 'selected_ids': selected, 'occurrence_ids': occurrences}


def transfer(client, body):
    result = client.post('/api/new-items/handoff', json=body)
    assert result.status_code == 200, result.text
    identity = result.json()['workspace_id']
    bundles = client.post('/api/workspaces/' + identity + '/bundles', json={})
    assert bundles.status_code == 200, bundles.text
    return client.get('/api/workspaces/' + identity).json(), bundles.json()['bundles']


def test_send_all_keeps_each_selected_occurrence_bound_to_its_id(client, test_database, monkeypatch):
    pending = preview(client, [('A', [['First', 123, 0, 0, 'GENERAL', False]]),
                               ('B', [['Second', 123, 0, 0, 'GENERAL', False]])], game=GAME)
    drafts = apply_preview(client, pending, ['A', 'B'], game=GAME)
    _, bundles = transfer(client, pass_drafts(client, test_database, monkeypatch, drafts))
    assert [(b['group'], [i['id'] for i in b['items']]) for b in bundles] == [('A', ['91']), ('B', ['92'])]


def test_converged_final_payloads_are_rejected_before_any_post(client, test_database, monkeypatch):
    first, _ = prepare(client)
    second, _ = prepare(client)
    _connect_aztek(client)
    observed = boundary(monkeypatch, test_database, first['id'], lose=True)
    response = run(client, first, second)
    assert response.status_code == 409, response.text
    assert not observed['posts'] and not observed['gets']


def legacy(client, sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in [['Event Name', 'Plan'], ['Reward: Group A'],
                    ['ItemKind', 'ItemOption', 'DurationIndex', 'Display Name', 'Amt'], *rows]:
            sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    response = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'event'},
                           files={'file': ('legacy.xlsx', stream.getvalue())})
    assert response.status_code == 200, response.text
    return apply_preview(client, response.json(), [s[0] for s in sheets], game=GAME)


@pytest.mark.parametrize('separate_sheets', [False, True])
def test_import_groups_belong_to_source_not_draft(client, test_database, monkeypatch, separate_sheets):
    sheets = [('A', [[123, 0, 0, 'First', 2], [123, 0, 0, 'Second', 3]])]
    if separate_sheets:
        sheets = [('A', [[123, 0, 0, 'Same', 2]]), ('B', [[123, 0, 0, 'Same', 3]])]
    drafts = legacy(client, sheets)
    workspace, bundles = transfer(client, pass_drafts(client, test_database, monkeypatch, drafts))
    assert len(bundles) == (2 if separate_sheets else 1), workspace
    assert sorted(i['qty'] for b in bundles for i in b['items']) == ['2', '3']


def test_saved_without_readback_blocks_next_scope_after_restart(client, test_database, monkeypatch):
    first, _ = prepare(client)
    second, _ = prepare(client, name='Next')
    with test_database.session() as db:
        job = db.get(Job, first['id'])
        job.status = 'saved'
        job.result = {'creation': {'state': 'saved', 'item_id': '91', 'source': 'created', 'evidence': {
            'detail_url': f'{ORIGIN}/combo/cabalm/shop/items/91'}}}
    _connect_aztek(client)
    observed = boundary(monkeypatch, test_database, second['id'], lose=True)
    with reopen(client, test_database) as restored:
        response = run(restored, second)
        assert response.status_code == 409, response.text
        row = {r['id']: r for r in restored.get('/api/new-items/history').json()['rows']}[first['id']]
        assert row['result']['creation']['item_id'] == '91'
        assert row['result']['recheck']['incomplete'] is True
        assert run(restored, first).status_code == 409
    assert not observed['posts'] and not observed['gets']


def test_dirty_selected_editor_cannot_freeze_old_values(client, operator_page):
    one_draft(client)
    page, requests = operator_page
    page.goto('http://tool.test/new-items')
    page.get_by_label('เลือกร่าง', exact=True).check()
    page.get_by_label('Name', exact=True).fill('Visible unsaved edit')
    page.get_by_role('button', name='ยืนยันร่างที่เลือก', exact=True).click()
    expect(page.get_by_role('dialog', name='ยืนยันร่าง', exact=True)).not_to_be_visible()
    expect(page.locator('#newNotice')).to_contain_text('บันทึกร่างก่อน')
    expect(page.get_by_label('Name', exact=True)).to_have_value('Visible unsaved edit')
    assert not any(path.endswith('/confirm') and method == 'POST' for method, path in requests)


def test_finder_stale_index_never_becomes_a_different_missing_item(client, member, test_database):
    from test_new_item_finder_entry_ui import source_workspace
    identity = source_workspace(member, test_database)
    snapshot = client.get('/api/workspaces/' + identity).json()
    body = {'workspace_id': identity, 'selected_missing_indexes': [0], 'game': GAME}
    if snapshot.get('new_item_snapshot'):
        body['snapshot_token'] = snapshot['new_item_snapshot']
    with test_database.session() as db:
        source = db.get(WorkspaceRecord, identity)
        source.results = [*source.results, {'aztek_id': '91', 'item_kind': '123',
                                         'item_option': '0', 'duration_index': '0'}]
    response = client.post('/api/new-items/from-finder', json=body)
    assert response.status_code == 409, response.text
    assert client.get('/api/new-items/drafts').json()['drafts'] == []


def test_finder_handoff_keeps_random_rate_tier_and_reference(client, member, test_database, monkeypatch):
    with test_database.session() as db:
        source = WorkspaceRepository(db).create(member.id, 'shop', 'random.xlsx')
        source.game = GAME
        source.criteria = [{'kind': '123', 'opt': '0', 'dur': '0', 'name': 'Synthetic name'}]
        reference = {'name': 'Synthetic name', 'kind': '123', 'option': '0', 'duration': '0',
                     'qty': '2', 'rate': '100', 'tier': 'Rare', 'filename': 'random.xlsx'}
        source.occurrences = [dict(source.criteria[0], amt='2', rate='100', tier='Rare',
            sources=['Box'], group_keys=['box'], source_sheet='A', position=0, document_reference=reference)]
        source.group_meta = {'box': {'group': 'Box', 'is_random': True}}
        identity = source.id
    snapshot = client.get('/api/workspaces/' + identity).json()
    body = {'workspace_id': identity, 'selected_missing_indexes': [0], 'game': GAME}
    if snapshot.get('new_item_snapshot'):
        body['snapshot_token'] = snapshot['new_item_snapshot']
    response = client.post('/api/new-items/from-finder', json=body)
    assert response.status_code == 200, response.text
    workspace, bundles = transfer(client, pass_drafts(client, test_database, monkeypatch, response.json()['drafts']))
    assert workspace['results'][0].get('rate') == '100'
    assert bundles[0]['is_random'] is True and bundles[0]['items'][0]['rate'] == '100'
    assert bundles[0]['document_reference']['items'][0]['tier'] == 'Rare'


def test_create_requires_current_search_before_browser(client, test_database, monkeypatch):
    entry, _ = prepare(client)
    with test_database.session() as db:
        job = db.get(Job, entry['id'])
        config = deepcopy(job.config)
        config.pop('search', None)
        job.config = config
    _connect_aztek(client)
    observed = boundary(monkeypatch, test_database, entry['id'], lose=True)
    response = run(client, entry)
    assert response.status_code == 409, response.text
    assert not observed['posts'] and not observed['gets']


def test_handoff_uses_source_row_order_not_newest_history_order(client, test_database, monkeypatch):
    pending = preview(client, [('A', [['First', 123, 0, 0, 'GENERAL', False],
                                     ['Second', 123, 0, 0, 'GENERAL', False]])], game=GAME)
    body = pass_drafts(client, test_database, monkeypatch, apply_preview(client, pending, ['A'], game=GAME))
    body['selected_ids'].reverse()
    workspace, _ = transfer(client, body)
    assert [r['item_name'] for r in workspace['results']] == ['First', 'Second']


@pytest.mark.parametrize('change', ['missing', 'other_search', 'other_payload'])
def test_create_rejects_search_not_bound_to_current_decision(client, test_database, monkeypatch, change):
    entry, _ = prepare(client)
    if change == 'missing':
        entry.pop('search_id')
    elif change == 'other_search':
        entry['search_id'] = 'not-the-shown-search'
    else:
        with test_database.session() as db:
            job = db.get(Job, entry['id'])
            config = deepcopy(job.config)
            config['search']['payload_hash'] = 'different-payload'
            job.config = config
    _connect_aztek(client)
    observed = boundary(monkeypatch, test_database, entry['id'], lose=True)
    assert run(client, entry).status_code == 409
    assert not observed['gets'] and not observed['posts']


@pytest.mark.parametrize('images', [False, True])
def test_edit_or_same_image_convergence_cannot_send_two_creates(client, test_database, monkeypatch, images):
    from test_new_item_create_flow import creation_entry
    from test_new_item_media import attach
    pending = preview(client, [('A', [['Same', 123, 0, None, 'GENERAL', False],
                                     ['Same', 123, 0, 0, 'GENERAL', False]])], game=GAME)
    drafts = apply_preview(client, pending, ['A'], game=GAME)
    entries = []
    for draft in drafts:
        edited = client.patch('/api/new-items/drafts/' + draft['draft_id'], json={
            'revision': draft['revision'], 'values': {'duration': '0'}}).json()
        if images:
            response = attach(client, edited)
            assert response.status_code == 200, response.text
            edited = response.json()
        frozen = confirm(client, edited, game=GAME)
        entries.append(creation_entry(client, edited['draft_id'], frozen))
    _connect_aztek(client)
    observed = boundary(monkeypatch, test_database, entries[0]['id'], lose=True)
    assert run(client, *entries).status_code == 409
    assert not observed['gets'] and not observed['posts']
    records = client.get('/api/new-items/history').json()['rows']
    assert len(records) == 2 and all(row['state'] == 'confirmed' for row in records)
    assert sum(len(row['frozen']['occurrences']) for row in records) == 2


def test_finder_missing_snapshot_is_rejected_even_when_index_in_range(client, member, test_database):
    from test_new_item_finder_entry_ui import source_workspace
    identity = source_workspace(member, test_database)
    response = client.post('/api/new-items/from-finder', json={
        'workspace_id': identity, 'selected_missing_indexes': [0], 'game': GAME})
    assert response.status_code == 409
    assert client.get('/api/new-items/drafts').json()['drafts'] == []


def test_direct_shop_import_keeps_random_zero_rates_through_send_all(client, test_database, monkeypatch):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Random Shop'
    for row in [['Product Name', 'Box'],
                ['ItemKind', 'ItemOption', 'DurationIndex', 'Item Name', 'Amt', 'Rate'],
                [123, 0, 0, 'First', 2, 0], [123, 0, 0, 'Second', 3, 100]]:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    response = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'shop'},
                           files={'file': ('random-shop.xlsx', stream.getvalue())})
    assert response.status_code == 200, response.text
    drafts = apply_preview(client, response.json(), ['Random Shop'], game=GAME)
    _, bundles = transfer(client, pass_drafts(client, test_database, monkeypatch, drafts))
    assert len(bundles) == 1 and bundles[0]['is_random'] is True
    assert [(i['id'], i['qty'], i['rate']) for i in bundles[0]['items']] == [('91', '2', '0'), ('92', '3', '100')]
    assert [r['rate'] for r in bundles[0]['document_reference']['items']] == ['0', '100']


def test_ui_requires_search_before_enabling_create_selection(client, operator_page):
    draft = one_draft(client, ['Unsearched', 123, 0, 0, 'GENERAL', False], game=GAME)
    confirm(client, draft, game=GAME)
    page, requests = operator_page
    page.goto('http://tool.test/new-items?game=CabalM%20TH')
    expect(page.get_by_role('checkbox', name='เลือกสร้างใหม่ Unsearched', exact=True)).to_be_disabled()
    expect(page.get_by_role('button', name='สร้าง Item ที่เลือก', exact=True)).to_be_disabled()
    assert not any(method == 'POST' and path.endswith('/run') for method, path in requests)
