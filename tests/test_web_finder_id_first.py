"""Real Item Finder runner/API; only external Aztek browser traffic is replaced."""
import asyncio
import json
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import urlsplit

import pytest
from playwright.async_api import async_playwright
from playwright.sync_api import sync_playwright, expect
from openpyxl import Workbook

from web import search_runner
from web.aztek_locator import AztekLocatorError, require_control
from test_web_search_auth import _connect_aztek, _drain


LIST_HTML = '''<main>
<form onsubmit="event.preventDefault()">
  <input type="text" placeholder="ค้นหา Item...">
  <button type="submit">ค้นหา</button>
</form>
<input id="filter-option" placeholder="ItemOption">
<input id="filter-duration" placeholder="DurationIndex">
<table><thead><tr><th>Aztek Item Id</th><th>ชื่อ</th><th>ItemKind</th>
<th>ItemOption</th><th>DurationIndex</th></tr></thead>
<tbody><tr><td>91</td><td>Water</td><td>12345</td><td>0</td><td>9</td></tr></tbody></table>
<div><button disabled>‹ Previous</button><input type="number" value="1">
<button disabled>Next ›</button></div>
</main>'''


def routed_browser(monkeypatch, html, *, detail=None, context_options=None):
    """Launch a real browser; unexpected requests abort rather than reach Aztek."""
    requests = []
    unexpected = []

    async def serve(route):
        request = route.request
        requests.append((request.method, request.url))
        path = urlsplit(request.url).path
        if request.method != 'GET':
            unexpected.append(request.url)
            await route.abort()
        elif path.endswith('/shop/items'):
            await route.fulfill(content_type='text/html; charset=utf-8', body=html)
        elif detail is not None and path.endswith('/shop/items/91'):
            await route.fulfill(content_type='text/html; charset=utf-8', body=detail)
        else:
            unexpected.append(request.url)
            await route.abort()

    class RoutedPlaywright:
        async def __aenter__(self):
            self.driver = await async_playwright().start()

            async def launch(**kwargs):
                browser = await self.driver.chromium.launch(**kwargs)

                async def new_context(**context_kwargs):
                    context = await browser.new_context(**{**context_kwargs, **(context_options or {})})
                    await context.route('**/*', serve)
                    return context

                return SimpleNamespace(new_context=new_context, close=browser.close)

            return SimpleNamespace(chromium=SimpleNamespace(launch=launch))

        async def __aexit__(self, *_args):
            await self.driver.stop()

    monkeypatch.setattr(search_runner, 'async_playwright', RoutedPlaywright)
    return requests, unexpected


def run_finder(*, deep=False):
    logs = []
    results = []
    finder = search_runner.HeadlessFinder(
        lambda message, level: logs.append((message, level)),
        results.append, lambda *_args: None)
    data = {
        'game': 'CabalM TH',
        'url': 'https://aztek-tools-v2.combo-interactive.com/combo/cabalm/shop/items',
        'headless': True, 'deep': deep,
        'read_desc': deep,
        'web': 'yes', 'img': 'any', 'qty_val': '5',
        'trade': 'yes', 'drill': 'no', 'crit_val': '20',
        'multi': [{'kind': '12345', 'opt': '0', 'dur': '9'}],
    }
    asyncio.run(finder.run(data, {'cookies': [], 'origins': []}))
    return results, logs


def test_ambiguous_search_stops_instead_of_reading_unfiltered_items(monkeypatch):
    # Existing first-match probes silently choose one of two search forms.
    html = LIST_HTML.replace('<input id="filter-option"', '''
    <form><input type="text" placeholder="ค้นหา Item...">
    <button type="submit">ค้นหา</button></form><input id="filter-option"''')
    requests, unexpected = routed_browser(monkeypatch, html)
    with pytest.raises(AztekLocatorError, match='ค้นหา'):
        run_finder()
    assert unexpected == []
    assert all(method == 'GET' for method, _url in requests)


DETAIL_HTML = '''<main><form>
<input id="item-name" value="Water"><input id="item-game-item-id" value="12345">
<input id="item-option" value="0"><input id="duration-index" value="9">
<label>เปิดใช้งานการแสดงผลบนเว็บ<input type="checkbox"></label>
<label>จำนวน<input value="999"></label><label>คริติคอล<input value="999"></label>
<textarea name="detail">wrong description</textarea>
<section><button id="display-option-enable" role="switch" aria-checked="true">Changed label</button></section>
<button id="game-tradeable" role="checkbox" aria-checked="true">Changed trade label</button>
<button id="displayOptions-is_slotable" role="checkbox" aria-checked="false">Changed drill label</button>
<input id="displayOptions-quantity" value="5"><input id="displayOptions-critical" value="20">
<textarea id="item-detail">correct description</textarea>
</form></main>'''


def test_deep_search_reads_stable_ids_before_misleading_labels(monkeypatch):
    requests, unexpected = routed_browser(monkeypatch, LIST_HTML, detail=DETAIL_HTML)
    results, _logs = run_finder(deep=True)
    assert [row['aztek_id'] for row in results] == ['91']
    assert results[0]['_desc'] == 'correct description'
    assert unexpected == []
    assert all(method == 'GET' for method, _url in requests)


def test_column_reorder_still_reads_the_correct_item_parameters(monkeypatch):
    html = LIST_HTML.replace(
        '<th>Aztek Item Id</th><th>ชื่อ</th><th>ItemKind</th>\n<th>ItemOption</th><th>DurationIndex</th>',
        '<th>ชื่อ</th><th>DurationIndex</th><th>ItemOption</th><th>Aztek Item Id</th><th>ItemKind</th>')
    html = html.replace(
        '<td>91</td><td>Water</td><td>12345</td><td>0</td><td>9</td>',
        '<td>Water</td><td>9</td><td>0</td><td>91</td><td>12345</td>')
    _requests, unexpected = routed_browser(monkeypatch, html)
    results, _logs = run_finder()
    assert [(r['aztek_id'], r['item_kind'], r['item_option'], r['duration_index'])
            for r in results] == [('91', '12345', '0', '9')]
    assert unexpected == []


@pytest.mark.parametrize('broken', ['missing', 'duplicate', 'wrong_role', 'missing_state', 'wrong_input_type'])
def test_broken_detail_contract_stops_instead_of_becoming_not_found(monkeypatch, broken):
    if broken == 'missing':
        detail = DETAIL_HTML.replace('id="display-option-enable"', 'id="changed"')
    elif broken == 'duplicate':
        detail = DETAIL_HTML.replace('</form>',
            '<button hidden id="display-option-enable" role="switch" aria-checked="false"></button></form>')
    elif broken == 'wrong_role':
        detail = DETAIL_HTML.replace('role="switch"', 'role="checkbox"')
    elif broken == 'missing_state':
        detail = DETAIL_HTML.replace('id="display-option-enable" role="switch" aria-checked="true"',
                                     'id="display-option-enable" role="switch"')
    else:
        detail = DETAIL_HTML.replace('id="displayOptions-quantity"',
                                     'type="checkbox" id="displayOptions-quantity"')
    requests, unexpected = routed_browser(monkeypatch, LIST_HTML, detail=detail)
    with pytest.raises(AztekLocatorError, match='แสดงผลบนเว็บ|จำนวน'):
        run_finder(deep=True)
    assert unexpected == []
    assert all(method == 'GET' for method, _url in requests)


def import_workspace(client):
    workbook = Workbook()
    workbook.active.append(['ItemKind', 'itemOption', 'durationIndex', 'ItemName'])
    workbook.active.append(['12345', '0', '9', 'Water'])
    output = BytesIO()
    workbook.save(output)
    response = client.post('/api/import-template', files={
        'file': ('fixture.xlsx', output.getvalue())}, data={'mode': 'event'})
    assert response.status_code == 200, response.text
    return response.json()['workspace_id']


def test_api_explains_locator_failure_and_preserves_imported_draft(
        client, anonymous_client, monkeypatch):
    _connect_aztek(client, anonymous_client)
    workspace_id = import_workspace(client)
    html = LIST_HTML.replace('<input id="filter-option"', '''
        <form><input placeholder="ค้นหา Item..."><button type="submit">ค้นหา</button></form>
        <input id="filter-option"''')
    requests, unexpected = routed_browser(monkeypatch, html)
    with client.websocket_connect('/ws/search') as ws:
        ws.send_json({'workspace_id': workspace_id, 'game': 'CabalM TH', 'web_mode': 'any'})
        messages = _drain(ws)
    errors = [m for m in messages if m.get('type') == 'error']
    assert any(m.get('code') == 'aztek_locator_changed' and 'ค้นหา' in m.get('msg', '')
               for m in errors), messages
    view = client.get('/api/workspaces/' + workspace_id).json()
    assert view['count'] == 1
    assert view['items'][0]['kind'] == '12345'
    assert view['result_count'] == 0
    assert client.get('/api/aztek/status').json()['status'] == 'active'
    assert unexpected == []
    assert all(method == 'GET' for method, _url in requests)


def test_pagination_ignores_outside_next_and_filters_after_all_pages(monkeypatch):
    html = LIST_HTML.replace('<td>0</td>', '<td>1</td>')
    html = html.replace('<button disabled>Next ›</button>', '''
    <button onclick="document.querySelector('main tbody').innerHTML=
      '<tr><td>92</td><td>Water page 2</td><td>12345</td><td>0</td><td>9</td></tr>';
      this.disabled=true">Next ›</button>''')
    html = '<button onclick="window.wrongNext=true">Next ›</button>' + html
    _requests, unexpected = routed_browser(monkeypatch, html)
    results, _logs = run_finder()
    assert [row['aztek_id'] for row in results] == ['92']
    assert unexpected == []


def test_detail_from_different_item_is_not_reported_as_a_match(monkeypatch):
    detail = DETAIL_HTML.replace('id="item-game-item-id" value="12345"',
                                 'id="item-game-item-id" value="99999"')
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML, detail=detail)
    with pytest.raises(AztekLocatorError, match='ItemKind'):
        run_finder(deep=True)
    assert unexpected == []


def test_missing_detail_kind_is_a_contract_error_not_an_empty_search(monkeypatch):
    detail = DETAIL_HTML.replace('id="item-game-item-id" value="12345"', 'id="changed" value=""')
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML, detail=detail)
    with pytest.raises(AztekLocatorError, match='ItemKind'):
        run_finder(deep=True)
    assert unexpected == []


@pytest.mark.parametrize('game', ['CabalM TH', 'CabalM SEA', 'CabalPC TH', 'CabalPC SEA'])
def test_api_search_persists_correct_item_for_each_verified_game(
        client, anonymous_client, monkeypatch, game):
    _connect_aztek(client, anonymous_client)
    workspace_id = import_workspace(client)
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML)
    with client.websocket_connect('/ws/search') as ws:
        ws.send_json({'workspace_id': workspace_id, 'game': game, 'web_mode': 'any'})
        messages = _drain(ws)
    assert not [m for m in messages if m.get('type') == 'error'], messages
    view = client.get('/api/workspaces/' + workspace_id).json()
    assert view['result_count'] == 1
    assert [(r['aztek_id'], r['item_kind'], r['item_option'], r['duration_index'], r['game'])
            for r in view['results']] == [('91', '12345', '0', '9', game)]
    assert view['results'][0]['document_reference']['kind'] == '12345'
    assert unexpected == []


def test_detail_waits_for_lazy_stable_switch_and_ignores_outside_duplicate(monkeypatch):
    switch = '<button id="display-option-enable" role="switch" aria-checked="true">Changed label</button>'
    detail = DETAIL_HTML.replace(switch, '''<span id="mount"></span><script>
      setTimeout(() => document.getElementById('mount').outerHTML =
        '<button id="display-option-enable" role="switch" aria-checked="true">Late label</button>', 1500);
    </script>''')
    detail = '<button hidden id="display-option-enable" role="switch" aria-checked="false"></button>' + detail
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML, detail=detail)
    results, _logs = run_finder(deep=True)
    assert [r['aztek_id'] for r in results] == ['91']
    assert unexpected == []


def test_loading_placeholder_is_not_an_empty_completed_search(monkeypatch):
    row = '<tr><td>91</td><td>Water</td><td>12345</td><td>0</td><td>9</td></tr>'
    html = LIST_HTML.replace(row, '<tr><td colspan="5">Loading...</td></tr>')
    html = html.replace('onsubmit="event.preventDefault()"',
        "onsubmit=\"event.preventDefault();setTimeout(() => document.querySelector('tbody').innerHTML="
        "'<tr><td>91</td><td>Water</td><td>12345</td><td>0</td><td>9</td></tr>',2800)\"")
    _requests, unexpected = routed_browser(monkeypatch, html)
    results, _logs = run_finder()
    assert [row['aztek_id'] for row in results] == ['91']
    assert unexpected == []


def test_verified_empty_state_remains_not_found_not_a_locator_failure(monkeypatch):
    table = LIST_HTML[LIST_HTML.index('<table>'):LIST_HTML.index('</table>') + len('</table>')]
    html = LIST_HTML.replace(table, '<p>ไม่พบ Item ที่ตรงกับเงื่อนไข</p>')
    _requests, unexpected = routed_browser(monkeypatch, html)
    results, logs = run_finder()
    assert results == []
    assert any('ไม่พบ item ที่ตรงเงื่อนไขเลย' in message for message, _level in logs)
    assert unexpected == []


def test_unknown_empty_table_stops_instead_of_certifying_not_found(monkeypatch):
    row = '<tr><td>91</td><td>Water</td><td>12345</td><td>0</td><td>9</td></tr>'
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML.replace(row, ''))
    with pytest.raises(AztekLocatorError, match='ตาราง Item'):
        run_finder()
    assert unexpected == []


def test_search_wrong_input_type_stops_before_submit(monkeypatch):
    html = LIST_HTML.replace('type="text" placeholder="ค้นหา Item..."',
                             'type="number" placeholder="ค้นหา Item..."')
    _requests, unexpected = routed_browser(monkeypatch, html)
    with pytest.raises(AztekLocatorError, match='ช่องค้นหา Item'):
        run_finder()
    assert unexpected == []


def test_duplicate_preview_stops_instead_of_using_any_image(monkeypatch):
    detail = DETAIL_HTML.replace('</form>',
        '<img alt="Preview" src="data:,first"><img alt="Preview" src="data:,second"></form>')
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML, detail=detail)
    with pytest.raises(AztekLocatorError, match='Preview'):
        run_finder(deep=True)
    assert unexpected == []


def test_literal_punctuation_id_does_not_match_css_class_lookalike():
    # One focused browser-level contract test for the shared lookup. Finder's
    # verified IDs themselves have no dots; other flows must test their own IDs.
    async def check_literal():
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            context = await browser.new_context()
            await context.route('**/*', lambda route: route.abort())
            page = await context.new_page()
            await page.set_content('''<main><input id="displayOptions" class="quantity" value="wrong">
                <section><input id="displayOptions.quantity" value="correct"></section></main>''')
            field = await require_control(page.locator('main'), 'จำนวน',
                element_id='displayOptions.quantity', tags=('input',), input_types=('text',))
            assert await field.input_value() == 'correct'
            await context.close()
            await browser.close()
    asyncio.run(check_literal())


def test_real_runner_keeps_shared_occurrences_and_selected_bundle_scope(
        client, anonymous_client, monkeypatch):
    _connect_aztek(client, anonymous_client)
    workbook = Workbook()
    for index, amount in enumerate((2, 7)):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = 'Reward ' + str(index + 1)
        sheet.append(['Activity ' + str(index + 1)])
        sheet.append(['ItemKind', 'itemOption', 'durationIndex', 'ItemName', 'Amt'])
        sheet.append(['12345', '0', '9', 'Water', amount])
    output = BytesIO()
    workbook.save(output)
    imported = client.post('/api/import-plan', data={'mode': 'event'},
        files={'file': ('shared.xlsx', output.getvalue())})
    assert imported.status_code == 200, imported.text
    pending = imported.json()
    applied = client.post('/api/import-plan/apply', json={
        'pending_id': pending['pending_id'], 'selected_sheets': ['Reward 1', 'Reward 2']})
    assert applied.status_code == 200, applied.text
    workspace_id = pending['workspace_id']
    _requests, unexpected = routed_browser(monkeypatch, LIST_HTML)
    with client.websocket_connect('/ws/search') as ws:
        ws.send_json({'workspace_id': workspace_id, 'game': 'CabalM TH', 'web_mode': 'any'})
        messages = _drain(ws)
    assert not [m for m in messages if m.get('type') == 'error'], messages
    results = client.get('/api/workspaces/' + workspace_id).json()['results']
    assert [row['aztek_id'] for row in results] == ['91', '91']
    assert [row['amt'] for row in results] == ['2', '7']
    assert [row['document_reference']['sheet'] for row in results] == ['Reward 1', 'Reward 2']
    assert results[0]['group_keys'] != results[1]['group_keys']
    selected = client.post('/api/workspaces/' + workspace_id + '/bundles',
                           json={'selected_indexes': [1]}).json()['bundles']
    assert len(selected) == 1
    assert selected[0]['group_key'] == results[1]['group_keys'][0]
    assert [(item['id'], item['qty']) for item in selected[0]['items']] == [('91', '7')]
    all_bundles = client.post('/api/workspaces/' + workspace_id + '/bundles',
                             json={'selected_indexes': []}).json()['bundles']
    assert [bundle['items'][0]['qty'] for bundle in all_bundles] == ['2', '7']
    assert unexpected == []


@pytest.mark.parametrize('width', [1280, 390])
def test_real_api_failure_notice_survives_done_in_the_ui(
        client, anonymous_client, monkeypatch, width):
    _connect_aztek(client, anonymous_client)
    workspace_id = import_workspace(client)
    html = LIST_HTML.replace('<input id="filter-option"', '''
        <form><input placeholder="ค้นหา Item..."><button type="submit">ค้นหา</button></form>
        <input id="filter-option"''')
    _requests, unexpected = routed_browser(monkeypatch, html)
    with client.websocket_connect('/ws/search') as ws:
        ws.send_json({'workspace_id': workspace_id, 'game': 'CabalM TH', 'web_mode': 'any'})
        messages = _drain(ws)
    assert [m for m in messages if m.get('type') == 'error'], messages

    # Replay the real coordinator's messages only at the browser WebSocket
    # transport boundary; HTML, workspace loading and notice logic stay real.
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(viewport={'width': width, 'height': 900})
        context.route('**/*', lambda route: route.abort())

        def serve_app(route):
            request = route.request
            url = urlsplit(request.url)
            response = client.request(request.method, url.path +
                ('?' + url.query if url.query else ''), content=request.post_data_buffer,
                headers={'Content-Type': request.headers.get('content-type', 'application/json')})
            route.fulfill(status=response.status_code, body=response.content,
                          content_type=response.headers.get('content-type', 'text/plain'))

        context.route('http://tool.test/**', serve_app)
        context.add_init_script("localStorage.setItem('afc.workspaceId', %s);" % json.dumps(workspace_id))
        context.add_init_script('''window.WebSocket = class {
            constructor() { setTimeout(() => this.onopen?.(), 0); }
            send() { setTimeout(() => {
                for (const message of %s) this.onmessage?.({data:JSON.stringify(message)});
                this.onclose?.({code:1000}); window.failureReplayFinished = true;
            }, 0); }
        };''' % json.dumps(messages, ensure_ascii=False))
        page = context.new_page()
        page.goto('http://tool.test/', wait_until='domcontentloaded')
        button = page.locator('#btnSearch')
        expect(button).to_be_enabled()
        button.focus()
        button.press('Enter')
        page.wait_for_function('window.failureReplayFinished === true')
        notice = page.locator('#finderNotice')
        expect(notice).to_have_attribute('role', 'alert')
        expect(notice).to_contain_text('ฟอร์มค้นหา Item')
        expect(page.locator('#progressText')).to_contain_text('ไม่สำเร็จ')
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        context.close()
        browser.close()
    assert unexpected == []
