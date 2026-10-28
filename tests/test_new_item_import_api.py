"""New Item import exercises real HTTP/storage; no Aztek connection is needed."""
from hashlib import sha256
from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest

GAME = 'CabalPC TH'
HEADERS = ['Name', 'ItemKind', 'Option', 'Duration', 'Item Type', 'Tradeable', 'Description']


def workbook_bytes(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def preview(client, sheets, *, game=GAME):
    content = workbook_bytes(sheets)
    result = client.post('/api/new-items/import', data={'game': game, 'mode': 'template'},
                         files={'file': ('items.xlsx', content)})
    assert result.status_code == 200, result.text
    return result.json()


def apply_preview(client, pending, selected, *, game=GAME):
    result = client.post('/api/new-items/import/apply', json={
        'pending_id': pending['pending_id'], 'selected_sheets': selected, 'game': game})
    assert result.status_code == 200, result.text
    return result.json()['drafts']


def one_draft(client, row=None, *, game=GAME):
    row = row or ['Potion x2', 123, 0, 9, 'GENERAL', False, 'note']
    return apply_preview(client, preview(client, [('รายการจริง', [row])], game=game),
                         ['รายการจริง'], game=game)[0]


def test_sheet_selection_keeps_names_and_does_not_change_input(client, application):
    content = workbook_bytes([('ไทย จริง', [['Potion x2', 123, 0, 0, 'GENERAL', False]]),
                              ('Do not select', [['Other', 777, 1, 9, 'GENERAL', True]])])
    before = sha256(content).hexdigest()
    response = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'template'},
                           files={'file': ('original.xlsx', content)})
    assert response.status_code == 200, response.text
    pending = response.json()
    assert [s['name'] for s in pending['sheets']] == ['ไทย จริง', 'Do not select']
    assert client.get('/api/new-items/drafts').json()['drafts'] == []
    drafts = apply_preview(client, pending, ['ไทย จริง'])
    assert len(drafts) == 1
    assert drafts[0]['values']['name'] == 'Potion x2'
    assert drafts[0]['values']['option'] == '0'
    assert drafts[0]['values']['duration'] == '0'
    assert drafts[0]['values']['tradeable'] is False
    assert drafts[0]['occurrences'][0]['filename'] == 'original.xlsx'
    assert drafts[0]['occurrences'][0]['sheet'] == 'ไทย จริง'
    assert drafts[0]['occurrences'][0]['row'] == 2
    assert sha256(content).hexdigest() == before
    assert not application.state.browser_gate.busy


def test_template_download_roundtrip_is_a_preview_not_create(client):
    template = client.get('/api/new-items/template')
    assert template.status_code == 200
    book = load_workbook(BytesIO(template.content))
    assert [cell.value for cell in book.active[1]] == HEADERS
    book.active.append(['New', 9, 0, 0, 'GENERAL', False, ''])
    stream = BytesIO()
    book.save(stream)
    book.close()
    response = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'template'},
                           files={'file': ('template.xlsx', stream.getvalue())})
    assert response.status_code == 200, response.text
    assert sum(s['row_count'] for s in response.json()['sheets']) == 1
    assert client.get('/api/new-items/drafts').json()['drafts'] == []


@pytest.mark.parametrize('selected', [[], ['Missing'], ['A', 'A']])
def test_invalid_sheet_selection_never_adds_drafts(client, selected):
    pending = preview(client, [('A', [['One', 1, 0, 0, 'GENERAL', False]])])
    response = client.post('/api/new-items/import/apply', json={
        'pending_id': pending['pending_id'], 'selected_sheets': selected, 'game': GAME})
    assert response.status_code == 400
    assert client.get('/api/new-items/drafts').json()['drafts'] == []


def test_import_preview_and_drafts_are_owner_scoped(client, client_for, other_member):
    pending = preview(client, [('A', [['Private', 1, 0, 0, 'GENERAL', False]])])
    other = client_for(other_member)
    denied = other.post('/api/new-items/import/apply', json={
        'pending_id': pending['pending_id'], 'selected_sheets': ['A'], 'game': GAME})
    assert denied.status_code == 404
    draft = apply_preview(client, pending, ['A'])[0]
    assert other.get('/api/new-items/drafts').json()['drafts'] == []
    assert other.patch('/api/new-items/drafts/' + draft['draft_id'],
                       json={'revision': draft['revision'], 'values': {'name': 'stolen'}}).status_code == 404


def test_game_is_fixed_at_preview_not_changed_during_apply(client):
    pending = preview(client, [('A', [['One', 1, 0, 0, 'GENERAL', False]])])
    result = client.post('/api/new-items/import/apply', json={
        'pending_id': pending['pending_id'], 'selected_sheets': ['A'], 'game': 'CabalM TH'})
    assert result.status_code == 400
    assert client.get('/api/new-items/drafts').json()['drafts'] == []


def test_missing_duration_is_not_zero(client):
    draft = one_draft(client, ['Missing', 123, 0, None, 'GENERAL', False])
    assert draft['values']['option'] == '0'
    assert draft['values']['duration'] is None
    assert draft['ready'] is False
    result = client.post('/api/new-items/confirm', json={
        'selected': [{'id': draft['draft_id'], 'revision': draft['revision']}], 'game': GAME})
    assert result.status_code == 400


def test_unauthenticated_new_item_routes_are_denied(anonymous_client):
    assert anonymous_client.get('/api/new-items/template').status_code == 401
    assert anonymous_client.get('/api/new-items/drafts').status_code == 401


def test_legacy_event_plan_uses_display_name_not_sheet_name(client):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet is not the Item name'
    for row in [['Event Name', 'Plan'], ['Reward: Group A'],
                ['ItemKind', 'ItemOption', 'DurationIndex', 'Display Name', 'Amt'],
                [123, 0, 9, 'Potion x20', 2]]:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    result = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'event'},
                         files={'file': ('plan.xlsx', stream.getvalue())})
    assert result.status_code == 200, result.text
    draft = apply_preview(client, result.json(), [sheet.title])[0]
    assert draft['values']['name'] == 'Potion x20'
    assert draft['values']['tradeable'] is False
    assert draft['ready'] is False
    assert draft['occurrences'][0]['amount'] == '2'


def test_from_finder_preserves_only_selected_missing_occurrences(client, member, test_database):
    from web.workspaces import WorkspaceRepository
    with test_database.session() as db:
        workspace = WorkspaceRepository(db).create(member.id, 'event', 'finder.xlsx')
        workspace.game = GAME
        workspace.criteria = [
            {'kind': '123', 'opt': '0', 'dur': '9', 'name': 'Missing'},
            {'kind': '777', 'opt': '0', 'dur': '9', 'name': 'Found'},
            {'kind': '999', 'opt': '0', 'dur': '9', 'name': 'Unselected'}]
        workspace.occurrences = [
            {'kind': '123', 'opt': '0', 'dur': '9', 'name': 'Missing',
             'amt': '2', 'source_sheet': 'A', 'source_row': 2, 'group_keys': ['g1']},
            {'kind': '123', 'opt': '0', 'dur': '9', 'name': 'Missing',
             'amt': '3', 'source_sheet': 'B', 'source_row': 8, 'group_keys': ['g2']}]
        workspace.results = [{'item_kind': '777', 'item_option': '0', 'duration_index': '9',
                              'aztek_id': '10'}]
        workspace_id = workspace.id
    result = client.post('/api/new-items/from-finder', json={
        'workspace_id': workspace_id, 'selected_missing_indexes': [0], 'game': GAME,
        'snapshot_token': client.get('/api/workspaces/' + workspace_id).json()['new_item_snapshot']})
    assert result.status_code == 200, result.text
    drafts = result.json()['drafts']
    assert len(drafts) == 1
    assert drafts[0]['values']['name'] == 'Missing'
    assert [o['amount'] for o in drafts[0]['occurrences']] == ['2', '3']
    assert [o['group_key'] for o in drafts[0]['occurrences']] == ['g1', 'g2']
    assert all(o['workspace_id'] == workspace_id for o in drafts[0]['occurrences'])


def test_same_import_preview_is_applied_once_under_concurrent_requests(client):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    pending = preview(client, [('A', [['Once', 123, 0, 9, 'GENERAL', False]])])
    barrier = Barrier(2)
    def apply():
        barrier.wait()
        return client.post('/api/new-items/import/apply', json={
            'pending_id': pending['pending_id'], 'selected_sheets': ['A'], 'game': GAME})
    with ThreadPoolExecutor(2) as pool:
        responses = list(pool.map(lambda _: apply(), range(2)))
    assert sorted(r.status_code for r in responses) == [200, 409]
    assert len(client.get('/api/new-items/drafts').json()['drafts']) == 1


def test_legacy_missing_option_duration_stays_missing(client):
    book = Workbook()
    sheet = book.active
    for row in [['Event Name', 'Plan'], ['Reward: A'],
                ['ItemKind', 'ItemOption', 'DurationIndex', 'Display Name', 'Amt'],
                [123, None, None, 'Missing values', 0]]:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    result = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'event'},
                         files={'file': ('missing.xlsx', stream.getvalue())})
    assert result.status_code == 200
    draft = apply_preview(client, result.json(), [sheet.title])[0]
    assert draft['values']['option'] is None
    assert draft['values']['duration'] is None
    assert draft['occurrences'][0]['amount'] == '0'
    assert draft['ready'] is False


def test_new_item_import_uses_existing_workbook_limit_not_small_json_limit(client):
    from zipfile import ZipFile, ZIP_STORED
    stream = BytesIO(workbook_bytes([('A', [['Large file', 123, 0, 0, 'GENERAL', False]])]))
    with ZipFile(stream, 'a') as archive:
        archive.writestr('ignored-padding.bin', b'x' * (2 * 1024 * 1024 + 1), compress_type=ZIP_STORED)
    response = client.post('/api/new-items/import', data={'game': GAME, 'mode': 'template'},
                           files={'file': ('large.xlsx', stream.getvalue())})
    assert response.status_code == 200, response.text
    assert response.json()['sheets'][0]['row_count'] == 1


def test_unreadable_numeric_cell_remains_an_editable_draft(client):
    from datetime import datetime
    pending = preview(client, [('A', [['Date in Kind', datetime(2026, 10, 25), 0, 0, 'GENERAL', False]])])
    draft = apply_preview(client, pending, ['A'])[0]
    assert draft['ready'] is False
    assert draft['values']['kind'] != '0'
    assert '2026' in draft['values']['kind']
    repaired = client.patch('/api/new-items/drafts/' + draft['draft_id'], json={
        'revision': draft['revision'], 'values': {'kind': '123'}})
    assert repaired.status_code == 200
    assert repaired.json()['occurrences'][0]['original']['kind'] == draft['values']['kind']
