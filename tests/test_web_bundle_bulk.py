from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook


TEXT = ('กล่องสุ่มกิจกรรม\tRANDOM\n200479\t10\tCommon\t60\n'
        '200480\t3\tRare\t30\n200481\t1\tEpic\t10\n\n'
        'แพ็กแจกของ\tFIXED\n200479\t5\tCommon\nCREDIT: Alz\t100000\tRare\n')


def paste(client, text):
    return client.post('/api/bundles/import-text', json={'text': text})


def test_bulk_paste_keeps_types_amounts_rates_and_currency(client):
    response = paste(client, TEXT)
    assert response.status_code == 200
    sheet = response.json()['sheets'][0]
    assert sheet['errors'] == []
    first, second = sheet['bundles']
    assert first['name'] == 'กล่องสุ่มกิจกรรม'
    assert first['type'] == 'RANDOM'
    assert [(i['id'], i['qty'], i['tier'], i['rate']) for i in first['items']] == [
        ('200479', '10', 'Common', '60'), ('200480', '3', 'Rare', '30'),
        ('200481', '1', 'Epic', '10')]
    assert second['type'] == 'FIXED'
    assert second['items'][0]['id'] == '200479'
    assert second['rewards'][0]['type'] == 'CREDIT'
    assert second['rewards'][0]['value'] == 'Alz'
    assert second['rewards'][0]['qty'] == '100000'
    assert second['rewards'][0]['tier'] == 'Rare'


def test_excel_blocks_and_paste_produce_identical_bundles(client):
    book = Workbook()
    for line in TEXT.splitlines():
        book.active.append(line.split('\t'))
    out = BytesIO()
    book.save(out)
    book.close()
    excel = client.post('/api/bundles/import', files={'file': ('blocks.xlsx', out.getvalue())})
    assert excel.status_code == 200
    assert excel.json()['sheets'][0]['bundles'] == paste(client, TEXT).json()['sheets'][0]['bundles']


def test_extra_excel_column_is_not_silently_discarded(client):
    book = Workbook()
    book.active.append(['Pack','FIXED'])
    book.active.append([11, 1, 'Common', None, None, 'unexpected'])
    out = BytesIO()
    book.save(out)
    book.close()
    response = client.post('/api/bundles/import', files={'file':('extra.xlsx', out.getvalue())})
    sheet = response.json()['sheets'][0]
    assert sheet['bundles'] == []
    assert sheet['errors'][0]['row'] == 2


def test_template_currency_placeholder_is_not_ready_to_create(client):
    sheet = paste(client, 'Pack\tFIXED\nCREDIT: [ชื่อ Currency]\t1000\tCommon').json()['sheets'][0]
    assert sheet['bundles'] == []
    assert sheet['errors'][0]['row'] == 2


@pytest.mark.parametrize('text, row, message', [
    ('Pack\tRANDOM\n11\t1\tCommon', 2, 'เรท'),
    ('Pack\tRANDOM\n11\t1\tCommon\t50', 1, '100'),
    ('Pack\tRANDOM\n11\t1\tCommon\t0', 2, 'เรท'),
    ('Pack\tRANDOM\n11\t1\tCommon\t101', 2, 'เรท'),
    ('Pack\tFIXED\n11\t1\tCommon\t50', 2, 'RANDOM'),
    ('11\t1\tCommon', 1, 'บันเดิล'),
    ('Pack\tFIXED', 1, 'ไม่มี'),
    ('Pack\tFIXED\nCREDIT: Alz\t2\tWrong', 2, 'Tier'),
    ('Pack\tFIXED\nAlz\t2\tCommon', 2, 'Item ID'),
    ('Pack\tFIXED\n11\t1\tCommon\n11\t2\tRare', 3, 'ซ้ำ'),
    ('Pack\tFIXED\n11\t1\tCommon\t\tunexpected', 2, 'คอลัมน์'),
])
def test_bulk_invalid_data_never_adds_partial_bundles(client, text, row, message):
    response = paste(client, text)
    assert response.status_code == 200
    sheet = response.json()['sheets'][0]
    assert sheet['bundles'] == []
    assert any(e['row'] == row and message in e['message'] for e in sheet['errors'])


def test_paste_requires_auth(anonymous_client):
    assert paste(anonymous_client, TEXT).status_code == 401


def test_reward_rate_and_tier_survive_run_validation():
    from web.app import _clean_rewards
    assert _clean_rewards([{'type':'CREDIT','value':'Alz','qty':'2','tier':'Epic','rate':'25'}]) == [
        {'type':'CREDIT','value':'Alz','qty':'2','tier':'Epic','rate':'25'}]


def test_reward_tier_and_rate_are_filled_after_successful_item_rows():
    import asyncio
    from test_web_bundle import _builder, _async, FakePage
    builder = _builder()
    captured = {}
    async def qty_tier(page, entries):
        captured['entries'] = entries
        return True
    async def rates(page, entries):
        captured['rates'] = entries
        return True
    outcomes = iter([False, True])
    async def add_item(*args):
        return next(outcomes)
    builder._fill_header = lambda *args: _async(True)
    builder._add_item = add_item
    builder._add_reward = lambda *args: _async(True)
    builder._fill_qty_tier = qty_tier
    builder._fill_rates = rates
    builder._fill_blank_tiers = lambda *args: _async(0)
    item = {'id':'22','qty':'5','tier':'Rare','rate':'75'}
    reward = {'type':'CREDIT','value':'Alz','qty':'1000','tier':'Epic','rate':'25'}
    result = asyncio.run(builder._fill_form(FakePage(), 'Pack', 'RANDOM', True,
                        [{'id':'11'}, item], [reward]))
    assert tuple(result) == (1, 1)
    assert captured['entries'] == [item, reward]
    assert captured['rates'] == [item, reward]


def test_template_download_uses_copyable_block_layout(client):
    response = client.get('/api/bundles/template')
    book = load_workbook(BytesIO(response.content))
    sheet = book.active
    assert sheet['B1'].value == 'RANDOM'
    assert [sheet.cell(2, i).value for i in (2,3,4)] == [10, 'Common', 60]
    # Placeholders must force the operator to supply real Aztek IDs.
    for row, item_id in ((2, '11'), (3, '22'), (4, '33'), (7, '11')):
        sheet.cell(row, 1, item_id)
    sheet['A8'] = 'CREDIT: Alz'
    out = BytesIO()
    book.save(out)
    book.close()
    imported = client.post('/api/bundles/import', files={'file': ('filled.xlsx', out.getvalue())})
    assert imported.json()['sheets'][0]['errors'] == []
    assert len(imported.json()['sheets'][0]['bundles']) == 2
