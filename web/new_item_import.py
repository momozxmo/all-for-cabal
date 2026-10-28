"""New Item workbook previews; never modifies the source or contacts Aztek."""
from io import BytesIO
from copy import deepcopy

from openpyxl import Workbook, load_workbook

from web.new_item_domain import new_draft, normalize_values, text

HEADERS = ['Name', 'ItemKind', 'Option', 'Duration', 'Item Type', 'Tradeable', 'Description']
HEADER_FIELDS = dict(zip(HEADERS, ('name', 'kind', 'option', 'duration',
                                  'item_type', 'tradeable', 'description')))


def new_item_template_bytes() -> bytes:
    book = Workbook()
    book.active.title = 'New Items'
    book.active.append(HEADERS)
    book.active.freeze_panes = 'A2'
    for column in 'ABCDEFG':
        book.active.column_dimensions[column].width = 24
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def read_new_item_workbook(path: str, mode: str) -> dict:
    if mode == 'template':
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in book:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                header = [str(v or '').strip() for v in rows[0]]
                if not all(name in header for name in HEADERS):
                    raise ValueError('Sheet ' + sheet.title + ' ไม่มีหัวตาราง New Item ครบ')
                if len(header) != len(set(header)):
                    raise ValueError('หัวตารางซ้ำใน Sheet ' + sheet.title)
                fields = {field: header.index(name) for name, field in HEADER_FIELDS.items()}
                items = []
                for number, row in enumerate(rows[1:], 2):
                    if not any(v is not None and v != '' for v in row):
                        continue
                    values = {field: row[index] if index < len(row) else None
                              for field, index in fields.items()}
                    items.append({'values': normalize_values(values), 'row': number, 'amount': None,
                                  'group_key': '', 'position': number - 2})
                sheets.append({'name': sheet.title, 'hidden': sheet.sheet_state != 'visible',
                               'rows': items})
            return {'sheets': sheets, 'warnings': []}
        finally:
            book.close()
    from web.item_service import parser_for_mode, stamp_sheet_rows
    parsed, skipped = parser_for_mode(mode)(path)
    sheets = []
    for name, rows in parsed:
        items = []
        for position, row in enumerate(stamp_sheet_rows(name, rows)):
            values = {target: row.get(source) for target, source in
                      [('name', 'name'), ('kind', 'kind'), ('option', 'opt'), ('duration', 'dur')]}
            # Itemmove/trade from legacy plans has ambiguous meaning: never map it to Tradeable.
            items.append({'values': values, 'row': row.get('source_row'),
                          'amount': text(row.get('amt')), 'position': position,
                          'group_key': (row.get('group_keys') or [None])[0]
                          or str((row.get('group_meta') or {}).get('group_key') or row.get('group') or ''),
                          'downstream': {key: deepcopy(row[key]) for key in
                                         ('rate', 'tier', 'document_reference', 'group_meta') if key in row}})
        sheets.append({'name': name, 'hidden': False, 'rows': items})
    return {'sheets': sheets, 'warnings': [str(v) for v in skipped]}


def drafts_from_sheets(data: dict, names: list[str], game: str, filename: str) -> list[dict]:
    available = {sheet['name']: sheet for sheet in data['sheets']}
    if not names or len(names) != len(set(names)) or any(n not in available for n in names):
        raise ValueError('กรุณาเลือก Sheet ที่มีอยู่และไม่ซ้ำ')
    drafts = []
    for name in names:
        for row in available[name]['rows']:
            drafts.append(new_draft(game, row['values'], {
                'filename': filename, 'sheet': name, 'row': row['row'],
                'position': row['position'], 'amount': row['amount'],
                'group_key': row['group_key'], 'workspace_id': None,
                'downstream': deepcopy(row.get('downstream') or {})}))
    return drafts
