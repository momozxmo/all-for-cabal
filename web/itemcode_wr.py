# -*- coding: utf-8 -*-
"""Read Mastercode WR workbooks into reviewable Item Code drafts.

The importer is deliberately separate from the existing Item Code plan reader.
It reads only the metadata around each Master Code Block and never retains the
uploaded workbook or contacts Aztek.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.utils.datetime import from_excel


MASTER_HEADER = re.compile(
    r'\bmaster\s*code\b.*?\bcode\s*(\d+)\b', re.IGNORECASE)
GAME_NAMES = {
    'CABAL M TH': 'CabalM TH',
    'CABALM TH': 'CabalM TH',
    'CABAL PC TH': 'CabalPC TH',
    'CABALPC TH': 'CabalPC TH',
}
GAME_SUFFIXES = {'CabalM TH': 'mth', 'CabalPC TH': 'pcth'}
MISSING_MASTER_CODE = '[ยังไม่มี Mastercode]'


def _text(value) -> str:
    if value is None:
        return ''
    return str(value).strip()


def _normalized(value) -> str:
    return re.sub(r'[\s:_-]+', '', _text(value)).casefold()


def _header(value):
    raw = _text(value)
    match = MASTER_HEADER.search(raw)
    if not match:
        return None
    base = re.sub(r'\s+', ' ', match.group(0)).strip()
    return base, match.group(1)


def _label_kind(value):
    label = _normalized(value)
    if 'codeexpiredate' in label:
        return 'date'
    if label == 'game':
        return 'game'
    if label == 'mastercode':
        return 'mastercode'
    if label.startswith('limit') and 'ใช้งาน' in label:
        return 'limit'
    return None


def _values_to_right(sheet, row, column, *, stop_column, limit=None):
    values = []
    for candidate in range(column + 1, stop_column + 1):
        value = sheet.cell(row=row, column=candidate).value
        if _text(value):
            values.append((candidate, value))
            if limit and len(values) >= limit:
                break
    return values


def _metadata(sheet, anchor_row, anchor_column):
    found = {}
    max_row = min(sheet.max_row, anchor_row + 7)
    max_column = min(sheet.max_column, anchor_column + 10)
    min_column = max(1, anchor_column - 1)
    for row in range(anchor_row + 1, max_row + 1):
        for column in range(min_column, max_column + 1):
            kind = _label_kind(sheet.cell(row=row, column=column).value)
            if not kind or kind in found:
                continue
            right = _values_to_right(
                sheet, row, column, stop_column=max_column,
                limit=2 if kind == 'mastercode' else 1)
            found[kind] = (row, column, right)
    return found


def _date_value(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch)
        except (TypeError, ValueError, OverflowError):
            return None
        return converted.date() if isinstance(converted, datetime) else converted
    text_value = _text(value)
    match = re.search(r'(?<!\d)(\d{1,2})/(\d{1,2})/(\d{2,4})(?!\d)',
                      text_value)
    if not match:
        return None
    day, month, year = (int(part) for part in match.groups())
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _positive_integer(value) -> str:
    if value is None or isinstance(value, bool):
        return ''
    try:
        number = Decimal(_text(value))
    except (InvalidOperation, ValueError):
        return ''
    if number <= 0 or number != number.to_integral_value():
        return ''
    return str(int(number))


def _game(value) -> str:
    return GAME_NAMES.get(re.sub(r'\s+', ' ', _text(value)).upper(), '')


def _slug(name, game):
    made = re.sub(r'-+', '-', re.sub(
        r'[^a-z0-9-]', '', re.sub(r'\s+', '-', name.lower()))).strip('-')
    suffix = GAME_SUFFIXES.get(game, '')
    if suffix:
        made = re.sub(r'-(?:mth|pcth)$', '', made)
        made = f'{made}-{suffix}' if made else suffix
    return made


def _reward(name, mastercode, bundle_id, usage_limit):
    return {
        'name_th': name,
        'name_en': name,
        'uses_per_user': '1',
        'limited': bool(usage_limit),
        'quantity': usage_limit,
        'remaining': usage_limit,
        'code_type': '1',
        'code_list': mastercode,
        'prefix': '',
        'num_codes': '',
        'bundle_id': bundle_id,
    }


def _draft(row):
    return {
        'name_th': row['name'],
        'name_en': row['name'],
        'slug': _slug(row['name'], row['game']),
        'uses_per_user': '1',
        'limited': bool(row['usage_limit']),
        'quantity': row['usage_limit'],
        'remaining': row['usage_limit'],
        'start_time': row['start_time'],
        'end_time': row['end_time'],
        'group': '',
        'source_mode': 'mastercode_wr',
        'rewards': [_reward(
            row['name'], row['mastercode'], row['bundle_id'],
            row['usage_limit'])],
        'wr_base_header': row['base_header'],
        'wr_sequence': row['daily_sequence'],
        'wr_game': row['game'],
        'wr_auto_names': True,
        'wr_auto_slug': True,
    }


def parse_workbook(path):
    """Return exact-sheet rows ordered by workbook position."""
    book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    rows = []
    try:
        for sheet_index, sheet in enumerate(book.worksheets):
            merged_right_edges = {
                (merged.min_row, merged.min_col): merged.max_col
                for merged in sheet.merged_cells.ranges
                if merged.min_row == merged.max_row
            }
            for row in sheet.iter_rows():
                for cell in row:
                    parsed_header = _header(cell.value)
                    if not parsed_header:
                        continue
                    base_header, code_number = parsed_header
                    metadata = _metadata(sheet, cell.row, cell.column)

                    def first(kind, preferred_offset=None):
                        label_row, label_column, values = metadata.get(
                            kind, (None, None, []))
                        if preferred_offset is not None and label_row is not None:
                            return sheet.cell(
                                row=label_row,
                                column=label_column + preferred_offset).value
                        return values[0][1] if values else None

                    raw_date = first('date', 2)
                    day = _date_value(raw_date, book.epoch)
                    raw_game = first('game', 2)
                    normalized_game = _game(raw_game)
                    master_meta = metadata.get(
                        'mastercode', (None, None, []))
                    master_row, master_label_column, _master_values = master_meta
                    mastercode = ''
                    bundle_value = None
                    if master_row is not None:
                        master_value_column = master_label_column + 2
                        mastercode = _text(sheet.cell(
                            row=master_row,
                            column=master_value_column).value)
                        bundle_column = merged_right_edges.get(
                            (master_row, master_value_column),
                            master_value_column) + 1
                        bundle_value = sheet.cell(
                            row=master_row, column=bundle_column).value
                    bundle_id = _positive_integer(bundle_value)
                    usage_limit = _positive_integer(first('limit'))
                    rows.append({
                        'sheet': sheet.title,
                        'sheet_index': sheet_index,
                        'row': cell.row,
                        'column': cell.column,
                        'block': f'Code {code_number}',
                        'base_header': base_header,
                        'raw_date': _text(raw_date),
                        'date': day.isoformat() if day else '',
                        'game': normalized_game,
                        'raw_game': _text(raw_game),
                        'mastercode': mastercode,
                        'bundle_id': bundle_id,
                        'usage_limit': usage_limit,
                    })
    finally:
        book.close()

    counters = defaultdict(int)
    contributing_sheets = defaultdict(set)
    for row in rows:
        key = (row['date'], row['game'])
        if all(key):
            contributing_sheets[key].add(row['sheet'])
            counters[key] += 1
            row['daily_sequence'] = counters[key]
        else:
            row['daily_sequence'] = 1

    overlap_keys = {
        key for key, sheets in contributing_sheets.items() if len(sheets) > 1}
    for row in rows:
        issues = []
        warnings = []
        if not row['date']:
            issues.append('ไม่พบหรืออ่าน CODE EXPIRE DATE ไม่ได้')
        if not row['game']:
            issues.append('ไม่รู้จัก Game: %s' % (row['raw_game'] or 'ไม่พบข้อมูล'))
        if not row['mastercode']:
            issues.append('ยังไม่มี Mastercode')
        if not row['bundle_id']:
            issues.append('Bundle ID ต้องเป็นจำนวนเต็มบวก')
        if not row['usage_limit']:
            issues.append('Usage Limit ต้องเป็นจำนวนเต็มบวก')
        key = (row['date'], row['game'])
        if key in overlap_keys:
            warnings.append('Same-Day Game Overlap: พบหลาย Sheet ในวันและเกมเดียวกัน')
        shown_mastercode = row['mastercode'] or MISSING_MASTER_CODE
        row['name'] = (
            f"{row['daily_sequence']}. {row['base_header']} - "
            f"{shown_mastercode}")
        if row['date']:
            day = date.fromisoformat(row['date'])
            row['start_time'] = datetime.combine(day, time.min).isoformat()
            row['end_time'] = datetime.combine(
                day, time(23, 59, 59)).isoformat()
        else:
            row['start_time'] = ''
            row['end_time'] = ''
        row['issues'] = issues
        row['warnings'] = warnings
        row['ready'] = not issues
        row['draft'] = _draft(row)

    by_sheet = []
    sheet_rows = defaultdict(list)
    for row in rows:
        sheet_rows[row['sheet']].append(row)
    for sheet in sorted(sheet_rows, key=lambda name: sheet_rows[name][0]['sheet_index']):
        by_sheet.append((sheet, sheet_rows[sheet]))
    return by_sheet


def prepare_preview(rows, *, pending_id, page_game):
    """Attach pending-specific identity and page-game selection state."""
    prepared = []
    for source in rows:
        row = dict(source)
        draft = dict(row['draft'])
        workbook_identity = row.get('workbook_fingerprint') or pending_id
        source_key = 'wr:%s:%s:%s:%s' % (
            workbook_identity, row['sheet_index'], row['row'], row['column'])
        draft['wr_source_id'] = source_key
        row['source_id'] = source_key
        row.pop('workbook_fingerprint', None)
        row.pop('wr_import_mode', None)
        row['draft'] = draft
        game_matches = bool(row['game']) and row['game'] == page_game
        row['selectable'] = game_matches
        row['selected'] = bool(row['ready'] and game_matches)
        if row['game'] and not game_matches:
            row['disabled_reason'] = (
                f"เกมในไฟล์เป็น {row['game']} แต่หน้าปัจจุบันเลือก {page_game}")
        elif not row['game']:
            row['disabled_reason'] = 'ไม่รู้จัก Game ในไฟล์'
        else:
            row['disabled_reason'] = ''
        prepared.append(row)
    return prepared
